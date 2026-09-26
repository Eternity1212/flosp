#!/usr/bin/env python3
"""把台账 Markdown 转成 Word，表格转成**原生 Word 表格**而不是一堆竖线文本。

只支持这批报告实际用到的语法：ATX 标题、GFM 表格、无序/任务列表、
引用块、围栏代码块、水平线，以及行内的 ``**粗体**`` / ``` `等宽` ```。
刻意不引入 Markdown 库——依赖越少，半年后越跑得起来。

用法::

    python scripts/md_to_docx.py reports/FedOSP_实验台账_终版_2026-09-27.md
    python scripts/md_to_docx.py in.md -o out.docx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import List, Optional

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
except ImportError:  # pragma: no cover
    raise SystemExit("需要 python-docx：pip install python-docx")

#: 行内标记：**粗体**、`等宽`。按出现顺序切分，保证嵌套顺序正确。
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
TABLE_SEP = re.compile(r"^\|[\s:|-]+\|$")


def add_inline(par, text: str) -> None:
    """把一行带行内标记的文本写进段落，保留粗体与等宽。"""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = par.add_run(piece[1:-1])
            run.font.name = "Menlo"
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            par.add_run(piece)


def split_row(line: str) -> List[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows: List[str]) -> None:
    """GFM 表格 → Word 原生表格。第一行当表头加粗。"""
    body = [r for r in rows if not TABLE_SEP.match(r.strip())]
    if not body:
        return
    cells = [split_row(r) for r in body]
    ncol = max(len(r) for r in cells)
    table = doc.add_table(rows=0, cols=ncol)
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(cells):
        wr = table.add_row().cells
        for j in range(ncol):
            par = wr[j].paragraphs[0]
            add_inline(par, row[j] if j < len(row) else "")
            if i == 0:
                for run in par.runs:
                    run.bold = True
    doc.add_paragraph()


def convert(md_path: Path, out_path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "PingFang SC"
    doc.styles["Normal"].font.size = Pt(10.5)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    buf_tbl: List[str] = []
    buf_code: List[str] = []
    in_code = False

    def flush_table() -> None:
        nonlocal buf_tbl
        if buf_tbl:
            add_table(doc, buf_tbl)
            buf_tbl = []

    for raw in lines:
        line = raw.rstrip()

        if line.strip().startswith("```"):
            if in_code:
                par = doc.add_paragraph()
                run = par.add_run("\n".join(buf_code))
                run.font.name = "Menlo"
                run.font.size = Pt(9)
                par.paragraph_format.left_indent = Pt(18)
                buf_code, in_code = [], False
            else:
                flush_table()
                in_code = True
            continue
        if in_code:
            buf_code.append(raw)
            continue

        if line.startswith("|"):
            buf_tbl.append(line)
            continue
        flush_table()

        if not line.strip():
            continue
        if line.startswith("---") and set(line.strip()) == {"-"}:
            doc.add_paragraph("─" * 42).alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)), 4))
            continue

        if line.startswith(">"):
            par = doc.add_paragraph()
            add_inline(par, line.lstrip("> ").rstrip())
            par.paragraph_format.left_indent = Pt(18)
            for run in par.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            continue

        m = re.match(r"^(\s*)[-*]\s+(\[[ x]\]\s*)?(.*)$", line)
        if m:
            indent, task, text = m.groups()
            par = doc.add_paragraph(style="List Bullet")
            if task:
                par.add_run("☑ " if "x" in task else "☐ ")
            add_inline(par, text)
            par.paragraph_format.left_indent = Pt(18 + 14 * (len(indent) // 2))
            continue

        add_inline(doc.add_paragraph(), line)

    flush_table()
    doc.save(out_path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=None)
    args = ap.parse_args()
    out: Optional[Path] = args.out or args.markdown.with_suffix(".docx")
    convert(args.markdown, out)
    n_tbl = sum(1 for ln in args.markdown.read_text(encoding="utf-8").splitlines()
                if TABLE_SEP.match(ln.strip()))
    print(f"✓ {out}（{out.stat().st_size / 1024:.0f} KB，{n_tbl} 张原生表格）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
